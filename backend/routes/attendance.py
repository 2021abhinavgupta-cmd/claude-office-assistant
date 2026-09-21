"""
Attendance & Employee Blueprint
Routes: /api/attendance/*, /api/employees/*
"""
import csv
import json
import logging
import re
from datetime import datetime, time as dt_time
from io import StringIO

from flask import Blueprint, Response, jsonify, request

from utils import (IST, _is_admin, _load_employees, _save_employees,
                   now_ist, today_ist)

logger = logging.getLogger(__name__)
attendance_bp = Blueprint("attendance", __name__)


def _verified_admin() -> bool:
    """Real session check for the two routes below (security warning #6's
    caveat): _is_admin(user_id) is bool(user_id) everywhere else in this app
    by design (gotcha #60, every logged-in employee is admin) -- but reading
    that user_id from an unauthenticated query string here meant ANY
    non-empty string passed, not just a real logged-in employee. This ties
    it to an actual valid session_token cookie instead."""
    from routes.auth import _verify_session
    token = request.cookies.get("session_token", "")
    user_id = _verify_session(token)
    return bool(user_id) and _is_admin(user_id)

# Work-day window for the daily_attendance summary (checkin_time/checkout_time
# -- what the Live Attendance dashboard cards and "Total" hours are computed
# from). A login/logout outside this window (e.g. checking Lumina at
# midnight) still happened and is still logged verbatim in the raw
# `attendance` audit table, but must not be allowed to stretch the day's
# recorded work window earlier/later than the business day actually runs.
WORK_START = dt_time(9, 0, 0)
WORK_END = dt_time(22, 30, 0)


def _clamp_work_time(hhmmss: str) -> str:
    try:
        h, m, s = (int(p) for p in hhmmss.split(":"))
        actual = dt_time(h, m, s)
    except Exception:
        return hhmmss
    if actual < WORK_START:
        return WORK_START.strftime("%H:%M:%S")
    if actual > WORK_END:
        return WORK_END.strftime("%H:%M:%S")
    return hhmmss


# ── DB helpers ────────────────────────────────────────────────────────────────

def _attendance_conn():
    from db import get_connection
    return get_connection()


def _attendance_payload():
    body = request.get_json(silent=True)
    if isinstance(body, dict):
        return body
    raw = request.get_data(cache=False, as_text=True) or ""
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except Exception:
        return {}


def _attendance_checkin(user_id: str):
    """First IST login of day wins. checkin_time is clamped to the work-day
    window (see WORK_START/WORK_END) so a login before 9am doesn't record the
    day as having started at 2am.

    Uses DO UPDATE ... WHERE checkin_time IS NULL rather than DO NOTHING:
    _attendance_ping() (the presence heartbeat) can create today's row first
    with checkin_time still NULL (it only ever writes last_seen_at). A plain
    DO NOTHING would then permanently skip setting checkin_time once that
    placeholder row exists, even on a genuine login -- this still preserves
    "first real checkin wins" (a NULL only gets filled once), it just also
    handles the row already existing from a ping with nothing in it yet."""
    d = today_ist()
    t = _clamp_work_time(now_ist())
    ts = datetime.now(IST).isoformat(timespec="seconds")
    conn = _attendance_conn()
    with conn:
        cur = conn.cursor()
        cur.execute(
            """INSERT INTO daily_attendance (user_id, date, checkin_time)
               VALUES (?, ?, ?)
               ON CONFLICT(user_id, date) DO UPDATE SET checkin_time = excluded.checkin_time
               WHERE daily_attendance.checkin_time IS NULL""",
            (user_id, d, t),
        )
        if cur.rowcount > 0:
            conn.execute(
                "INSERT INTO attendance (user_id, action, timestamp) VALUES (?, 'in', ?)",
                (user_id, ts),
            )
        cur.execute(
            "SELECT checkin_time FROM daily_attendance WHERE user_id=? AND date=?",
            (user_id, d),
        )
        row = cur.fetchone()
        stored_checkin = row[0] if row else t
    conn.close()
    return d, stored_checkin


def _attendance_ping(user_id: str) -> str:
    """Lightweight presence heartbeat -- called every 60s while a Lumina page
    is open (see auth.js). Only writes last_seen_at, nothing else. Real
    checkout is inferred later by sweep_stale_checkouts() (task_scheduler.py)
    from staleness of this timestamp, NOT from any unload/pagehide event --
    that approach was tried and reverted (see CLAUDE.md gotcha #70/#71: a
    pagehide-triggered checkout write fires on every internal page
    navigation, not just a real tab close, and the resulting write volume
    was implicated in a production DB hang)."""
    d = today_ist()
    t = now_ist()
    conn = _attendance_conn()
    with conn:
        conn.execute(
            """INSERT INTO daily_attendance (user_id, date, last_seen_at)
               VALUES (?, ?, ?)
               ON CONFLICT(user_id, date) DO UPDATE SET last_seen_at = excluded.last_seen_at""",
            (user_id, d, t),
        )
    conn.close()
    return t


def _attendance_checkout(user_id: str):
    """Always updates checkout_time to latest IST logout (UPSERT).
    checkout_time is clamped to the work-day window (see WORK_START/WORK_END)
    so a logout after 10:30pm doesn't stretch the recorded work day later."""
    d = today_ist()
    t = _clamp_work_time(now_ist())
    ts = datetime.now(IST).isoformat(timespec="seconds")
    conn = _attendance_conn()
    with conn:
        cur = conn.cursor()
        cur.execute(
            """INSERT INTO daily_attendance (user_id, date, checkout_time)
               VALUES (?, ?, ?)
               ON CONFLICT(user_id, date) DO UPDATE SET checkout_time = excluded.checkout_time""",
            (user_id, d, t),
        )
        conn.execute(
            "INSERT INTO attendance (user_id, action, timestamp) VALUES (?, 'out', ?)",
            (user_id, ts),
        )
        cur.execute(
            "SELECT checkout_time FROM daily_attendance WHERE user_id=? AND date=?",
            (user_id, d),
        )
        row = cur.fetchone()
        stored_checkout = row[0] if row else t
    conn.close()
    return d, stored_checkout


# ── Attendance routes ─────────────────────────────────────────────────────────

@attendance_bp.route("/api/attendance/checkin", methods=["POST"])
def attendance_checkin():
    body = _attendance_payload()
    user_id = str(body.get("user_id", "")).strip()
    if not user_id:
        return jsonify({"error": "user_id required"}), 400
    date_ist, checkin_time = _attendance_checkin(user_id)
    return jsonify({
        "success": True,
        "user_id": user_id,
        "date": date_ist,
        "checkin_time": checkin_time,
        "timezone": "IST",
    })


@attendance_bp.route("/api/attendance/checkout", methods=["POST"])
def attendance_checkout():
    body = _attendance_payload()
    user_id = str(body.get("user_id", "")).strip()
    if not user_id:
        return jsonify({"error": "user_id required"}), 400
    date_ist, checkout_time = _attendance_checkout(user_id)
    return jsonify({
        "success": True,
        "user_id": user_id,
        "date": date_ist,
        "checkout_time": checkout_time,
        "timezone": "IST",
    })


@attendance_bp.route("/api/attendance/ping", methods=["POST"])
def attendance_ping():
    body = _attendance_payload()
    user_id = str(body.get("user_id", "")).strip()
    if not user_id:
        return jsonify({"error": "user_id required"}), 400
    last_seen = _attendance_ping(user_id)
    return jsonify({"success": True, "last_seen_at": last_seen})


def sweep_stale_checkouts(stale_after_seconds: int = 240):
    """Called by task_scheduler.py every ~3min. For today's rows with a
    checkin but no checkout, whose last_seen_at is older than
    stale_after_seconds, mark checkout_time = last_seen_at -- i.e. the
    employee's tab went away (closed / crashed / lost network) and their
    heartbeat simply stopped. Explicit Logout (attendance_checkout above)
    already handles the instant case; this is the fallback for everyone
    who just closes the tab."""
    d = today_ist()
    now_dt = datetime.strptime(now_ist(), "%H:%M:%S")
    conn = _attendance_conn()
    cur = conn.cursor()
    cur.execute(
        """SELECT user_id, last_seen_at FROM daily_attendance
           WHERE date=? AND checkin_time IS NOT NULL
             AND checkout_time IS NULL AND last_seen_at IS NOT NULL""",
        (d,),
    )
    rows = cur.fetchall()
    swept = 0
    with conn:
        for user_id, last_seen_at in rows:
            try:
                last_seen_dt = datetime.strptime(last_seen_at, "%H:%M:%S")
            except ValueError:
                continue
            if (now_dt - last_seen_dt).total_seconds() >= stale_after_seconds:
                conn.execute(
                    "UPDATE daily_attendance SET checkout_time=? WHERE user_id=? AND date=?",
                    (_clamp_work_time(last_seen_at), user_id, d),
                )
                swept += 1
    conn.close()
    return swept


@attendance_bp.route("/api/attendance/summary", methods=["GET"])
def attendance_summary():
    user_id = request.args.get("user_id", "").strip()
    if not user_id:
        return jsonify({"error": "user_id required"}), 400
    conn = _attendance_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT date, checkin_time, checkout_time FROM daily_attendance WHERE user_id=? ORDER BY date DESC",
        (user_id,),
    )
    rows = cur.fetchall()
    conn.close()
    return jsonify({
        "user_id": user_id,
        "timezone": "IST",
        "records": [
            {"date": r[0], "checkin_time": r[1], "checkout_time": r[2]}
            for r in rows
        ],
    })


@attendance_bp.route("/api/attendance/today", methods=["GET"])
def attendance_today():
    # Optional ?date=YYYY-MM-DD to view a past day's attendance (Live
    # Attendance dashboard date picker); defaults to today when omitted,
    # same as before this param existed.
    date_ist = request.args.get("date", "").strip() or today_ist()
    conn = _attendance_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT user_id, checkin_time, checkout_time FROM daily_attendance WHERE date=?",
        (date_ist,),
    )
    rows = cur.fetchall()
    conn.close()
    records = [
        {"user_id": r[0], "date": date_ist, "checkin_time": r[1], "checkout_time": r[2]}
        for r in rows
    ]
    return jsonify({"date": date_ist, "timezone": "IST", "records": records})


@attendance_bp.route("/api/attendance/logs", methods=["GET"])
def attendance_logs():
    if not _verified_admin():
        return jsonify({"error": "Unauthorized"}), 403
    from db import get_connection
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, user_id, action, timestamp FROM attendance ORDER BY timestamp DESC LIMIT 500")
    logs = [{"id": r[0], "user_id": r[1], "action": r[2], "timestamp": r[3]} for r in cursor.fetchall()]
    conn.close()
    return jsonify({"logs": logs})


@attendance_bp.route("/api/attendance/export", methods=["GET"])
def attendance_export():
    if not _verified_admin():
        return "Unauthorized", 403

    # daily_attendance (not the raw attendance event log) is already one row
    # per user per day with checkin_time/checkout_time -- both stored as
    # plain HH:MM:SS IST (see now_ist()/_clamp_work_time() above), so no
    # timezone conversion is needed here.
    from db import get_connection
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT date, user_id, checkin_time, checkout_time FROM daily_attendance ORDER BY date DESC, user_id"
    )
    rows = cursor.fetchall()

    # Standup task counts per (user_id, date), same definitions the Velocity
    # chart uses (routes/ops.py::get_velocity) so the two stay consistent:
    # completed = marked done that day; carried = still-pending tasks that
    # were carried in from an earlier day (not a new task started that day).
    cursor.execute(
        """SELECT user_id, date,
                  SUM(CASE WHEN status='done' THEN 1 ELSE 0 END) AS completed,
                  SUM(CASE WHEN carried_from IS NOT NULL AND status='pending' THEN 1 ELSE 0 END) AS carried
           FROM standup_tasks
           WHERE status != 'deleted'
           GROUP BY user_id, date"""
    )
    task_counts = {(r[0], r[1]): (r[2] or 0, r[3] or 0) for r in cursor.fetchall()}
    conn.close()

    emp_map = {}
    try:
        emps = _load_employees().get("employees", [])
        for e in emps:
            emp_map[e["id"]] = e["name"]
            if e.get("whatsapp"):
                emp_map[e["whatsapp"]] = e["name"]
                emp_map[e["whatsapp"].replace('+', '')] = e["name"]
    except Exception:
        pass

    def format_user(uid):
        if uid in emp_map:
            return emp_map[uid]
        if re.match(r"^\+?\d{10,15}$", uid):
            return f"WhatsApp ({uid[-4:]})"
        return uid

    si = StringIO()
    cw = csv.writer(si)
    cw.writerow(["Date", "In", "Out", "Employee", "Tasks Completed", "Tasks Carried Forward"])
    for date, user_id, checkin_time, checkout_time in rows:
        completed, carried = task_counts.get((user_id, date), (0, 0))
        cw.writerow([date, checkin_time or "", checkout_time or "", format_user(user_id), completed, carried])

    return Response(
        si.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=claude_attendance_logs.csv"}
    )


# Weekday-only Full Day / Half Day / Leave threshold. Matches the >=8h
# wording the user asked for; anything shorter but with a real checkin is
# a Half Day, and a weekday with no checkin at all is a Leave.
_FULL_DAY_HOURS = 8.0
_INACTIVE_STATUSES = {"inactive", "disabled", "left", "removed", "archived", "former"}


def _hours_and_day_type(checkin, checkout, is_today):
    """(hours:float|None, label:str) for one (checkin_time, checkout_time)
    pair, both plain 'HH:MM:SS' IST or falsy. No midnight-rollover handling
    needed -- both are already clamped inside the same work-day window by
    _clamp_work_time() before they're ever stored."""
    if not checkin:
        return None, "Leave"
    if not checkout:
        return None, ("In Progress" if is_today else "Incomplete")
    try:
        h1, m1, s1 = (int(p) for p in checkin.split(":"))
        h2, m2, s2 = (int(p) for p in checkout.split(":"))
        hrs = ((h2 * 3600 + m2 * 60 + s2) - (h1 * 3600 + m1 * 60 + s1)) / 3600.0
    except Exception:
        return None, "Incomplete"
    if hrs < 0:
        return None, "Incomplete"
    return round(hrs, 2), ("Full Day" if hrs >= _FULL_DAY_HOURS else "Half Day")


def _parse_time_obj(hhmmss):
    """'HH:MM:SS' -> a real datetime.time object (so Excel treats it as an
    actual time value, sortable/filterable, not a left-aligned text
    string) -- None if missing or unparseable."""
    if not hhmmss:
        return None
    try:
        h, m, s = (int(p) for p in hhmmss.split(":"))
        return dt_time(h, m, s)
    except Exception:
        return None


def _attendance_credit(day_type):
    """Half Day leave best practice (ExcelDemy/Indzara/Clockify-style HR
    templates researched for this feature): a Half Day contributes 0.5 to
    an attendance-percentage numerator, a Full Day contributes 1.0, and
    Leave/Incomplete contribute 0. "In Progress" (today, still ongoing) is
    excluded from both the numerator and the denominator entirely --
    there's no final answer for it yet."""
    if day_type == "Full Day":
        return 1.0
    if day_type == "Half Day":
        return 0.5
    return 0.0


@attendance_bp.route("/api/attendance/export-sheets", methods=["GET"])
def attendance_export_sheets():
    """An .xlsx workbook (not a flat CSV, so it can hold multiple sheets):
    an "All" sheet (same combined view the plain CSV export has, plus Day
    Type), a "Summary" sheet with one row per employee (Full/Half/Leave
    counts, attendance %, total hours), and one additional sheet per
    active employee with every WEEKDAY from their own start date through
    today filled in as its own row -- a day with no checkin at all shows
    as Leave rather than just being absent from the list.

    Each employee's sheet starts from config/employees.json's optional
    "joined_date" (YYYY-MM-DD) if set; otherwise from THAT employee's own
    earliest daily_attendance record (not a company-wide date), so a
    recently-onboarded employee doesn't get backfilled with Leave for
    months before they ever used the app. Set "joined_date" by hand in
    employees.json for anyone whose real hire date predates their first
    login, to backfill Leave correctly from the actual join date instead.

    Dates/times are written as REAL Excel date/time values (DD/MM/YYYY,
    HH:MM display format -- matches the dd/mm/yyyy convention already
    used app-wide, see gotcha #84), not plain text, so they sort/filter
    correctly in Excel instead of alphabetically. Attendance % is a real
    percentage number too, not a formatted string. Every sheet also gets
    an auto-filter on its header row and thin borders on every data cell."""
    if not _verified_admin():
        return "Unauthorized", 403

    import io
    import re as _re
    from datetime import date as _date, timedelta as _timedelta
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from db import get_connection
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT date, user_id, checkin_time, checkout_time FROM daily_attendance ORDER BY date DESC, user_id"
    )
    rows = cursor.fetchall()
    cursor.execute(
        """SELECT user_id, date,
                  SUM(CASE WHEN status='done' THEN 1 ELSE 0 END) AS completed,
                  SUM(CASE WHEN carried_from IS NOT NULL AND status='pending' THEN 1 ELSE 0 END) AS carried
           FROM standup_tasks
           WHERE status != 'deleted'
           GROUP BY user_id, date"""
    )
    task_counts = {(r[0], r[1]): (r[2] or 0, r[3] or 0) for r in cursor.fetchall()}
    conn.close()

    emp_map = {}
    active_employees = []
    try:
        for e in _load_employees().get("employees", []):
            emp_map[e["id"]] = e["name"]
            if e.get("whatsapp"):
                emp_map[e["whatsapp"]] = e["name"]
                emp_map[e["whatsapp"].replace("+", "")] = e["name"]
            if str(e.get("status", "active")).strip().lower() not in _INACTIVE_STATUSES:
                active_employees.append(e)
    except Exception:
        logger.exception("attendance_export_sheets: employee load failed")

    def format_user(uid):
        if uid in emp_map:
            return emp_map[uid]
        if re.match(r"^\+?\d{10,15}$", uid):
            return f"WhatsApp ({uid[-4:]})"
        return uid

    by_user_date = {}
    earliest_by_user = {}
    for d, uid, cin, cout in rows:
        by_user_date[(uid, d)] = (cin, cout)
        if uid not in earliest_by_user or d < earliest_by_user[uid]:
            earliest_by_user[uid] = d

    today = today_ist()
    today_d = _date.fromisoformat(today)

    wb = Workbook()
    HEADER_FILL = PatternFill("solid", fgColor="1F2937")
    HEADER_FONT = Font(color="FFFFFF", bold=True)
    TITLE_FONT = Font(bold=True, size=13)
    META_FONT = Font(italic=True, color="4B5563")
    SUMMARY_LABEL_FONT = Font(bold=True)
    THIN = Side(style="thin", color="D1D5DB")
    THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    DATE_FMT = "DD/MM/YYYY"
    TIME_FMT = "HH:MM"
    DAY_TYPE_FILLS = {
        "Full Day": PatternFill("solid", fgColor="D1FAE5"),
        "Half Day": PatternFill("solid", fgColor="FEF3C7"),
        "Leave": PatternFill("solid", fgColor="FEE2E2"),
    }

    def style_header(ws, headers, row=1):
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col_idx, value=h)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center")
            c.border = THIN_BORDER
        # A plain coordinate STRING, not ws.cell(...).coordinate -- calling
        # .cell() to merely read a coordinate still reserves that cell in
        # openpyxl's internal sheet dimensions, which silently bumps
        # max_row and shifts every later ws.append() down by one row.
        ws.freeze_panes = f"A{row + 1}"

    def border_row(ws, row, ncols):
        for c in range(1, ncols + 1):
            ws.cell(row=row, column=c).border = THIN_BORDER

    def set_widths(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 1: "All" -- same combined view as the plain CSV export,
    # plus a Day Type column. Unlike the per-employee sheets below, this
    # one is NOT filled in with missing-day rows -- it's one row per real
    # daily_attendance record, same as before.
    ws_all = wb.active
    ws_all.title = "All"
    ws_all.sheet_properties.tabColor = "374151"
    style_header(ws_all, ["Date", "In", "Out", "Employee", "Tasks Completed",
                          "Tasks Carried Forward", "Day Type"])
    for d, uid, cin, cout in rows:
        completed, carried = task_counts.get((uid, d), (0, 0))
        _, day_type = _hours_and_day_type(cin, cout, d == today)
        cin_obj, cout_obj = _parse_time_obj(cin), _parse_time_obj(cout)
        ws_all.append([_date.fromisoformat(d), cin_obj, cout_obj, format_user(uid),
                       completed, carried, day_type])
        r = ws_all.max_row
        ws_all.cell(row=r, column=1).number_format = DATE_FMT
        if cin_obj is not None:
            ws_all.cell(row=r, column=2).number_format = TIME_FMT
        if cout_obj is not None:
            ws_all.cell(row=r, column=3).number_format = TIME_FMT
        fill = DAY_TYPE_FILLS.get(day_type)
        if fill:
            ws_all.cell(row=r, column=7).fill = fill
        border_row(ws_all, r, 7)
    if ws_all.max_row > 1:
        ws_all.auto_filter.ref = f"A1:G{ws_all.max_row}"
    set_widths(ws_all, [12, 10, 10, 18, 14, 18, 12])

    # ── Per-employee sheets: each starts from its own real start date
    # (joined_date override, else that employee's own earliest
    # daily_attendance record, else today for a brand-new nobody's-logged-
    # in-yet employee) through today, every WEEKDAY filled in as a row.
    used_titles = set()
    emp_summaries = []  # collected for the "Summary" overview sheet below
    for emp in active_employees:
        uid = emp["id"]
        name = emp.get("name") or uid
        joined = str(emp.get("joined_date") or "").strip()
        start = None
        if joined:
            try:
                start = _date.fromisoformat(joined)
            except ValueError:
                logger.warning("attendance_export_sheets: bad joined_date %r for %s", joined, uid)
        if start is None:
            own_earliest = earliest_by_user.get(uid)
            start = _date.fromisoformat(own_earliest) if own_earliest else today_d
        if start > today_d:
            start = today_d
        end = today_d

        title = _re.sub(r'[\[\]:\*\?/\\]', "", name)[:31] or uid
        base_title, n = title, 2
        while title in used_titles:
            title = f"{base_title[:28]}~{n}"
            n += 1
        used_titles.add(title)

        ws = wb.create_sheet(title=title)
        ws.cell(row=1, column=1, value=f"Attendance -- {name}").font = TITLE_FONT
        meta_bits = [b for b in [emp.get("role"), emp.get("department")] if b]
        meta = " | ".join(meta_bits + [
            f"Period: {start.strftime('%d/%m/%Y')} to {end.strftime('%d/%m/%Y')}"
        ])
        ws.cell(row=2, column=1, value=meta).font = META_FONT
        style_header(ws, ["Date", "Day", "In", "Out", "Hours Worked", "Day Type",
                          "Tasks Completed", "Tasks Carried Forward"], row=4)

        counts = {"Full Day": 0, "Half Day": 0, "Leave": 0, "Incomplete": 0, "In Progress": 0}
        total_hours = 0.0
        credit_total = 0.0
        d = start
        while d <= end:
            if d.weekday() < 5:  # Mon-Fri only -- weekends aren't a "Leave"
                dstr = d.isoformat()
                cin, cout = by_user_date.get((uid, dstr), (None, None))
                hrs, day_type = _hours_and_day_type(cin, cout, dstr == today)
                completed, carried = task_counts.get((uid, dstr), (0, 0))
                cin_obj, cout_obj = _parse_time_obj(cin), _parse_time_obj(cout)
                ws.append([d, d.strftime("%A"), cin_obj, cout_obj,
                          hrs if hrs is not None else None, day_type, completed, carried])
                r = ws.max_row
                ws.cell(row=r, column=1).number_format = DATE_FMT
                if cin_obj is not None:
                    ws.cell(row=r, column=3).number_format = TIME_FMT
                if cout_obj is not None:
                    ws.cell(row=r, column=4).number_format = TIME_FMT
                if hrs is not None:
                    ws.cell(row=r, column=5).number_format = "0.00"
                fill = DAY_TYPE_FILLS.get(day_type)
                if fill:
                    ws.cell(row=r, column=6).fill = fill
                border_row(ws, r, 8)
                counts[day_type] = counts.get(day_type, 0) + 1
                credit_total += _attendance_credit(day_type)
                if hrs is not None:
                    total_hours += hrs
            d += _timedelta(days=1)
        data_end_row = ws.max_row
        if data_end_row >= 5:
            ws.auto_filter.ref = f"A4:H{data_end_row}"
        set_widths(ws, [12, 12, 10, 10, 14, 12, 14, 18])

        considered = counts["Full Day"] + counts["Half Day"] + counts["Leave"] + counts["Incomplete"]
        pct = (credit_total / considered) if considered else None

        summary_row = data_end_row + 2
        ws.cell(row=summary_row, column=1, value="Summary").font = SUMMARY_LABEL_FONT
        summary_lines = [
            ("Full Days", counts["Full Day"]),
            ("Half Days", counts["Half Day"]),
            ("Leaves", counts["Leave"]),
            ("Incomplete (no checkout logged)", counts["Incomplete"]),
            ("Total Working Days Considered", considered),
            ("Total Hours Worked", round(total_hours, 1)),
            ("Attendance %", pct if pct is not None else "N/A"),
        ]
        for i, (label, val) in enumerate(summary_lines, start=1):
            ws.cell(row=summary_row + i, column=1, value=label).font = SUMMARY_LABEL_FONT
            val_cell = ws.cell(row=summary_row + i, column=2, value=val)
            if label == "Total Hours Worked":
                val_cell.number_format = "0.0"
            elif label == "Attendance %" and val != "N/A":
                val_cell.number_format = "0.0%"

        emp_summaries.append({
            "name": name, "start": start, "full": counts["Full Day"],
            "half": counts["Half Day"], "leave": counts["Leave"],
            "incomplete": counts["Incomplete"], "hours": round(total_hours, 1), "pct": pct,
        })

    # ── "Summary" overview sheet: one row per employee, placed right after
    # "All" so it's the second tab (before diving into individual sheets).
    ws_sum = wb.create_sheet(title="Summary", index=1)
    ws_sum.sheet_properties.tabColor = "D97706"
    ws_sum.cell(row=1, column=1, value="Attendance Summary").font = TITLE_FONT
    ws_sum.cell(row=2, column=1, value=f"As of {today_d.strftime('%d/%m/%Y')}").font = META_FONT
    ws_sum.cell(row=3, column=1, value="Legend:").font = META_FONT
    for j, (label, fill) in enumerate(DAY_TYPE_FILLS.items(), start=2):
        c = ws_sum.cell(row=3, column=j, value=label)
        c.fill = fill
        c.alignment = Alignment(horizontal="center")
    ws_sum.cell(
        row=4, column=1,
        value=("Full Day = worked 8+ hours  |  Half Day = worked under 8 hours "
               "(but checked in)  |  Leave = no check-in that weekday  |  "
               "Incomplete = checked in but no checkout was logged  |  "
               "In Progress = still checked in today, not final yet"),
    ).font = META_FONT
    style_header(ws_sum, ["Employee", "Period Start", "Full Days", "Half Days",
                          "Leaves", "Incomplete", "Attendance %", "Total Hours Worked"], row=6)
    for s in emp_summaries:
        ws_sum.append([s["name"], s["start"], s["full"], s["half"], s["leave"],
                      s["incomplete"], s["pct"] if s["pct"] is not None else "N/A", s["hours"]])
        r = ws_sum.max_row
        ws_sum.cell(row=r, column=2).number_format = DATE_FMT
        ws_sum.cell(row=r, column=8).number_format = "0.0"
        if s["pct"] is not None:
            ws_sum.cell(row=r, column=7).number_format = "0.0%"
        border_row(ws_sum, r, 8)
    if ws_sum.max_row >= 7:
        ws_sum.auto_filter.ref = f"A6:H{ws_sum.max_row}"
    set_widths(ws_sum, [18, 14, 10, 10, 10, 10, 14, 16])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment;filename=lumina_attendance.xlsx"},
    )


# ── Employee routes ───────────────────────────────────────────────────────────

@attendance_bp.route("/api/employees", methods=["GET"])
def get_employees():
    """Returns all employees with their current status."""
    data = _load_employees()
    return jsonify(data)


@attendance_bp.route("/api/employees/checkin", methods=["POST"])
def employee_checkin():
    """
    Records an employee check-in or check-out.
    Body: { emp_id OR whatsapp, action: 'in'|'out', notes? }
    """
    body     = request.get_json(silent=True) or {}
    emp_id   = body.get("emp_id", "").strip()
    whatsapp = body.get("whatsapp", "").strip()
    action   = body.get("action", "in").strip()
    notes    = body.get("notes", "")

    if not emp_id and not whatsapp:
        return jsonify({"error": "emp_id or whatsapp required"}), 400

    data = _load_employees()
    found = None
    for emp in data["employees"]:
        if emp_id and emp.get("id") == emp_id:
            found = emp
            break
        if whatsapp and emp.get("whatsapp") == whatsapp:
            found = emp
            break

    if not found:
        return jsonify({"error": "Employee not found"}), 404

    entry = {
        "timestamp": datetime.now(IST).isoformat(timespec="seconds"),
        "action":    action,
        "notes":     notes,
    }
    found.setdefault("checkins", []).append(entry)
    found["last_action"] = action
    found["last_seen"]   = entry["timestamp"]

    _save_employees(data)
    if action == "out":
        date_ist, time_ist = _attendance_checkout(found.get("id", ""))
    else:
        date_ist, time_ist = _attendance_checkin(found.get("id", ""))
    logger.info(f"Employee {found['name']} checked {action} at {entry['timestamp']}")

    return jsonify({
        "success":  True,
        "employee": found["name"],
        "action":   action,
        "time":     entry["timestamp"],
        "date":     date_ist,
        "time_ist": time_ist,
        "timezone": "IST",
    })


@attendance_bp.route("/api/employees/summary", methods=["GET"])
def employee_summary():
    """Returns today's attendance summary."""
    data  = _load_employees()
    today = today_ist()
    conn = _attendance_conn()
    cur = conn.cursor()
    cur.execute(
        "SELECT user_id, checkin_time, checkout_time FROM daily_attendance WHERE date=?",
        (today,),
    )
    attendance_map = {r[0]: {"checkin_time": r[1], "checkout_time": r[2]} for r in cur.fetchall()}
    conn.close()

    summary = []
    for emp in data["employees"]:
        daily = attendance_map.get(emp.get("id", ""), {})
        checkin_time = daily.get("checkin_time")
        checkout_time = daily.get("checkout_time")
        if checkin_time and not checkout_time:
            status = "in"
        elif checkin_time and checkout_time:
            status = "out"
        else:
            status = "not checked in"
        summary.append({
            "emp_id":        emp.get("id", ""),
            "name":          emp["name"],
            "role":          emp["role"],
            "department":    emp["department"],
            "status":        status,
            "checkin_time":  checkin_time,
            "checkout_time": checkout_time,
            "today_logs":    [],
        })

    return jsonify({"date": today, "employees": summary, "total": len(summary)})
