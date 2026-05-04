"""
File Processor — Extract content from uploaded files for Claude context.
Supports: Images (vision), PDF, DOCX, XLSX, plain text, code files.
"""
import base64, logging
from pathlib import Path

logger = logging.getLogger(__name__)
MAX_TEXT = 50_000  # max chars extracted from any document

TEXT_EXTS = {
    '.txt', '.md', '.py', '.js', '.ts', '.jsx', '.tsx', '.html', '.css',
    '.json', '.yaml', '.yml', '.xml', '.csv', '.sql', '.sh', '.go',
    '.java', '.c', '.cpp', '.cs', '.php', '.rb', '.swift', '.kt', '.rs',
    '.toml', '.ini', '.env', '.r', '.scala', '.dart',
}
IMAGE_MIMES = {'image/jpeg', 'image/png', 'image/gif', 'image/webp'}


def process_file(file_bytes: bytes, filename: str, mime_type: str = '') -> dict:
    """
    Extract usable content from a file.
    Returns one of:
      {type:'image',    media_type, data (base64), filename, size_bytes}
      {type:'document', content (str),              filename}
      {type:'error',    error,                       filename}
    """
    ext = Path(filename).suffix.lower()

    # ── Images → base64 for Claude vision ──────────────────────────────────
    if mime_type in IMAGE_MIMES or ext in {'.jpg', '.jpeg', '.png', '.gif', '.webp'}:
        mt = mime_type if mime_type in IMAGE_MIMES else f"image/{ext.lstrip('.')}"
        return {
            'type':       'image',
            'media_type': mt,
            'data':       base64.standard_b64encode(file_bytes).decode(),
            'filename':   filename,
            'size_bytes': len(file_bytes),
        }

    # ── PDF ─────────────────────────────────────────────────────────────────
    if ext == '.pdf' or 'pdf' in mime_type:
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(file_bytes))
            text   = "\n\n".join(p.extract_text() or '' for p in reader.pages)
            return {'type': 'document', 'content': text[:MAX_TEXT], 'filename': filename,
                    'pages': len(reader.pages)}
        except ImportError:
            return {'type': 'error', 'error': 'pypdf not installed (pip install pypdf)', 'filename': filename}
        except Exception as e:
            return {'type': 'error', 'error': f'PDF error: {e}', 'filename': filename}

    # ── Word DOCX ────────────────────────────────────────────────────────────
    if ext == '.docx' or 'wordprocessingml' in mime_type:
        try:
            from docx import Document
            import io
            doc  = Document(io.BytesIO(file_bytes))
            text = "\n".join(p.text for p in doc.paragraphs)
            return {'type': 'document', 'content': text[:MAX_TEXT], 'filename': filename}
        except ImportError:
            return {'type': 'error', 'error': 'python-docx not installed', 'filename': filename}
        except Exception as e:
            return {'type': 'error', 'error': f'DOCX error: {e}', 'filename': filename}

    # ── Excel ────────────────────────────────────────────────────────────────
    if ext in {'.xlsx', '.xls'} or 'spreadsheetml' in mime_type:
        try:
            import openpyxl, io
            wb    = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"=== Sheet: {ws.title} ===")
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(c or '') for c in row))
                if sum(len(l) for l in lines) > MAX_TEXT: break
            return {'type': 'document', 'content': "\n".join(lines)[:MAX_TEXT], 'filename': filename}
        except ImportError:
            return {'type': 'error', 'error': 'openpyxl not installed', 'filename': filename}
        except Exception as e:
            return {'type': 'error', 'error': f'Excel error: {e}', 'filename': filename}

    # ── Text / Code / CSV → UTF-8 decode ─────────────────────────────────────
    if ext in TEXT_EXTS or mime_type.startswith('text/'):
        try:
            return {'type': 'document',
                    'content': file_bytes.decode('utf-8', errors='replace')[:MAX_TEXT],
                    'filename': filename}
        except Exception as e:
            return {'type': 'error', 'error': str(e), 'filename': filename}

    # ── Last resort ───────────────────────────────────────────────────────────
    try:
        return {'type': 'document',
                'content': file_bytes.decode('utf-8', errors='replace')[:MAX_TEXT],
                'filename': filename}
    except Exception:
        return {'type': 'error', 'error': 'Unsupported file type', 'filename': filename}
