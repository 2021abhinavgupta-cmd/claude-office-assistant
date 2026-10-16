"""Daily-standup Excel workbook, laid out like the attendance export
(routes/attendance.py::attendance_export_sheets): an "All" sheet, a "Summary"
sheet with one row per person, and one sheet per person with every task plus
a Monthly Summary underneath.

standup_tasks re-inserts a still-open task as a NEW row every day it stays
pending (carried_from keeps the date it first went unfinished). So counting
rows would count one stuck task once per day. Everything that says "tasks"
below counts LINEAGES instead: (person, title, origin date) where origin is
carried_from or the row's own date. A lineage is Done if any of its rows is
done (the done row is only ever written on the day it was finished).
"""
from __future__ import annotations

import io
import json
import re
from datetime import date as _date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="2F6B45")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10.5)
_TITLE_FONT = Font(bold=True, size=13, color="1E4D2F")
_META_FONT = Font(italic=True, color="4B5563", size=9.5)
_BOLD = Font(bold=True)
_BANNER_FILL = PatternFill("solid", fgColor="1E4D2F")
_BANNER_FONT = Font(bold=True, size=16, color="FFFFFF")
_SUB_FILL = PatternFill("solid", fgColor="E3EFE7")
_ZEBRA = PatternFill("solid", fgColor="F4F9F5")
_TOTAL_FILL = PatternFill("solid", fgColor="DCEBE0")
_CARD_FILL = PatternFill("solid", fgColor="F4F9F5")
_STATUS_FILL = {
    "Done": PatternFill("solid", fgColor="D1FAE5"),
    "Pending": PatternFill("solid", fgColor="FEF3C7"),
    "Blocked": PatternFill("solid", fgColor="FEE2E2"),
    "Delegated": PatternFill("solid", fgColor="DBEAFE"),
    "Needs Approval": PatternFill("solid", fgColor="EDE9FE"),
}
_DATE_FMT = "DD/MM/YYYY"
_SKIP_STATUS = {"deleted"}

ALL_HEADERS = ["Date", "Day", "Employee", "Task", "Status", "Due Date", "Days Carried",
               "Blocker", "Subtasks", "Delegated From", "Delegated To", "On Board", "Added At"]
ALL_WIDTHS = [12, 11, 16, 48, 14, 12, 12, 32, 12, 16, 16, 10, 18]
EMP_HEADERS = ["Date", "Day", "Task", "Status", "Due Date", "Days Carried",
               "Blocker", "Subtasks", "Delegated From", "Delegated To", "On Board"]
EMP_WIDTHS = [12, 11, 48, 14, 12, 12, 32, 12, 16, 16, 10]


def _d(s):
    try:
        return _date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def _status_label(status, blocker):
    s = (status or "pending").strip().lower()
    if s == "done":
        return "Done"
    if s == "delegated":
        return "Delegated"
    if s in ("need_for_approval", "needs_approval"):
        return "Needs Approval"
    return "Blocked" if (blocker or "").strip() else "Pending"


def _subtasks(raw):
    try:
        items = json.loads(raw) if isinstance(raw, str) else (raw or [])
    except (TypeError, ValueError):
        items = []
    items = [i for i in items if isinstance(i, dict)]
    if not items:
        return ""
    return f"{sum(1 for i in items if i.get('done'))}/{len(items)}"


def _norm_row(r, name_of):
    """DB row -> plain dict, with the derived columns every sheet needs."""
    (uid, date, title, status, blocker, carried_from, created_at,
     notion_id, subtasks, d_to, d_from, due) = r
    day = _d(date)
    origin = _d(carried_from) or day
    carried_days = (day - origin).days if day and origin and day > origin else 0
    return {
        "uid": uid, "name": name_of(uid), "date": day, "title": (title or "").strip(),
        "status": _status_label(status, blocker), "due": _d(due), "carried_days": carried_days,
        "blocker": (blocker or "").strip(), "subtasks": _subtasks(subtasks),
        "from": name_of(d_from) if d_from and str(d_from).startswith("emp") else (d_from or ""),
        "to": name_of(d_to) if d_to and str(d_to).startswith("emp") else (d_to or ""),
        "notion_id": notion_id or "", "on_board": "Yes" if notion_id else "", "added": str(created_at or "")[:16],
        "lineage": (uid, (title or "").strip().lower(), origin),
    }


def _header(ws, headers, row):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = _HEADER_FONT, _HEADER_FILL, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _border(ws, row, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).border = _BORDER


def _style_status(ws, row, col, label):
    fill = _STATUS_FILL.get(label)
    if fill:
        ws.cell(row=row, column=col).fill = fill


def _wrap(ws, row, cols):
    for c in cols:
        ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="center")


def banner(ws, title, subtitle, ncols, note=""):
    """Rows 1-2 (and 3 for a note): a dark-green title band, a pale
    subtitle band, gridlines off -- the sheet's 'letterhead'."""
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font, c.fill = _BANNER_FONT, _BANNER_FILL
    c.alignment = Alignment(vertical="center", indent=1)
    for i in range(2, ncols + 1):
        ws.cell(row=1, column=i).fill = _BANNER_FILL
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font, s.fill = Font(bold=True, color="1E4D2F", size=10.5), _SUB_FILL
    s.alignment = Alignment(vertical="center", indent=1)
    for i in range(2, ncols + 1):
        ws.cell(row=2, column=i).fill = _SUB_FILL
    ws.row_dimensions[2].height = 22
    if note:
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
        n = ws.cell(row=3, column=1, value=note)
        n.font = _META_FONT
        n.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[3].height = 30


def zebra(ws, first, last, ncols):
    """Alternate-row tint, never over a cell that already has its own fill
    (status chips keep their colour)."""
    for r in range(first, last + 1):
        if (r - first) % 2 == 1:
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                if cell.fill is None or cell.fill.fill_type is None:
                    cell.fill = _ZEBRA


def fit_row(ws, row, texts_widths, line_h=14.5, min_h=20):
    """Row height that fits the longest wrapped cell: [(text, col_width)]."""
    lines = 1
    for text, width in texts_widths:
        if text:
            per_line = max(8, int(width * 1.1))
            n = sum(max(1, -(-len(part) // per_line)) for part in str(text).split("\n"))
            lines = max(lines, n)
    ws.row_dimensions[row].height = max(min_h, line_h * lines + 5)


def page_setup(ws, header_row=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "Lumina  |  Daily Standup  |  Page &P of &N"


def lineage_stats(rows):
    """{lineage: {"done": bool, "blocked": bool, "month": (y, m)}}"""
    out = {}
    handed_off = {r["lineage"] for r in rows if r["status"] == "Delegated"}
    owned = {r["lineage"] for r in rows if r["status"] != "Delegated"}
    for r in rows:
        key = r["lineage"]
        # A task handed to a teammate now belongs to THEM (their copy is its
        # own row on their list); counting it for the delegator too would
        # count the same work twice.
        if key in handed_off and key not in owned:
            continue
        s = out.setdefault(key, {"done": False, "blocked": False,
                                 "month": (key[2].year, key[2].month) if key[2] else None})
        if r["status"] == "Done":
            s["done"] = True
        if r["blocker"]:
            s["blocked"] = True
    return out


def _pct(done, total):
    return (done / total) if total else None


_GREEN = PatternFill("solid", fgColor="2F6B45")
_TRACK_STATUS = ["Closed", "Open", "Blocked", "Needs Approval"]
_TRACK_FILL = {
    "Closed": PatternFill("solid", fgColor="CFE2F3"),
    "Open": PatternFill("solid", fgColor="FEF3C7"),
    "Blocked": PatternFill("solid", fgColor="F8CBCB"),
    "Needs Approval": PatternFill("solid", fgColor="E4D7F5"),
}
TRACK_HEADERS = ["Sector", "Client", "Type", "Goals", "TAT", "Start", "Finish", "Status",
                 "Time Track", "Time Taken", "QR", "Blocker"]
TRACK_WIDTHS = [13, 18, 16, 50, 7, 12, 12, 15, 14, 11, 7, 30]


def tracker_rows(rows, notion_info):
    """One dict per TASK (lineage) -- the one-row-per-task tracker layout.
    A task handed to a teammate is skipped here; it appears on theirs."""
    groups = {}
    for r in rows:
        groups.setdefault(r["lineage"], []).append(r)
    out = []
    for key, rs in groups.items():
        if all(r["status"] == "Delegated" for r in rs):
            continue
        rs = [r for r in rs if r["status"] != "Delegated"] or rs
        rs.sort(key=lambda r: r["date"])
        last = rs[-1]
        origin = key[2] or rs[0]["date"]
        done = [r["date"] for r in rs if r["status"] == "Done"]
        if done:
            status = "Closed"
        elif last["status"] == "Needs Approval":
            status = "Needs Approval"
        elif last["blocker"]:
            status = "Blocked"
        else:
            status = "Open"
        nid = next((r["notion_id"] for r in rs if r["notion_id"]), "")
        info = notion_info.get(nid, {}) if nid else {}
        due = last["due"] or _d(info.get("due"))
        tat = ((due - origin).days + 1) if due and due >= origin else None
        out.append({
            "name": last["name"], "sector": info.get("sector", ""),
            "client": info.get("client", ""),
            "type": info.get("type") or ("Daily task" if not nid else "Other"),
            "goal": last["title"], "tat": tat, "start": origin,
            "finish": min(done) if done else None, "status": status,
            "blocker": last["blocker"],
        })
    out.sort(key=lambda t: (t["start"], t["goal"].lower()))
    return out


def write_tracker(ws, trows, first_row, with_employee=False):
    """Header at first_row, tasks below, live formulas for Time Taken and
    Time Track so HR can edit TAT / dates / status and see it update.
    Returns the last data row."""
    off = 1 if with_employee else 0
    headers = (["Employee"] if with_employee else []) + TRACK_HEADERS
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=first_row, column=i, value=h)
        c.font, c.fill, c.border = _HEADER_FONT, _GREEN, _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[first_row].height = 28

    def col(n):  # 1-based tracker column -> letter, shifted if Employee leads
        return get_column_letter(n + off)

    r = first_row
    for t in trows:
        r += 1
        vals = [t["sector"], t["client"], t["type"], t["goal"], t["tat"], t["start"],
                t["finish"], t["status"], None, None, None, t["blocker"]]
        if with_employee:
            vals = [t["name"]] + vals
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        ws[f"{col(10)}{r}"] = (f'=IF(AND({col(6)}{r}<>"",{col(7)}{r}<>""),'
                               f'{col(7)}{r}-{col(6)}{r}+1,"")')
        ws[f"{col(9)}{r}"] = (f'=IF(AND({col(8)}{r}="Closed",{col(10)}{r}<>""),IF({col(5)}{r}="","No TAT",'
                              f'IF({col(10)}{r}>{col(5)}{r},"Late",IF({col(10)}{r}={col(5)}{r},'
                              f'"On Time","Before Time"))),"Task Not Closed")')
        for n in (6, 7):
            ws[f"{col(n)}{r}"].number_format = _DATE_FMT
        ws[f"{col(10)}{r}"].number_format = "0"
        fill = _TRACK_FILL.get(t["status"])
        if fill:
            ws[f"{col(8)}{r}"].fill = fill
        for n in (5, 6, 7, 8, 9, 10, 11):
            ws[f"{col(n)}{r}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"{col(9)}{r}"].font = _BOLD
        ws[f"{col(4)}{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws[f"{col(12)}{r}"].alignment = Alignment(wrap_text=True, vertical="center")
        for n in (1, 2, 3):
            ws[f"{col(n)}{r}"].alignment = Alignment(vertical="center")
        _border(ws, r, len(headers))
        fit_row(ws, r, [(t["goal"], TRACK_WIDTHS[3]), (t["blocker"], TRACK_WIDTHS[11])])
    last = r
    zebra(ws, first_row + 1, last, len(headers))
    if last > first_row:
        rng = f"{col(9)}{first_row + 1}:{col(9)}{last}"
        for text, color in (("Late", "C0392B"), ("On Time", "1E8449"), ("Before Time", "1F618D")):
            ws.conditional_formatting.add(
                rng, FormulaRule(formula=[f'{col(9)}{first_row + 1}="{text}"'],
                                 font=Font(bold=True, color=color)))
        dv = DataValidation(type="list", formula1='"' + ",".join(_TRACK_STATUS) + '"',
                            allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col(8)}{first_row + 1}:{col(8)}{last}")
        ws.auto_filter.ref = f"A{first_row}:{get_column_letter(len(headers))}{last}"
    for i, w in enumerate(([16] if with_employee else []) + TRACK_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return last


def build_workbook(raw_rows, emp_map, active_employees, today: _date,
                   notion_info=None) -> bytes:
    """raw_rows: SELECT user_id, date, title, status, blocker, carried_from,
    created_at, notion_id, subtasks, delegated_to, delegated_from, due_date
    FROM standup_tasks, already filtered to the date range wanted."""
    notion_info = notion_info or {}

    def name_of(uid):
        return emp_map.get(uid) or str(uid)

    rows = [_norm_row(r, name_of) for r in raw_rows
            if str(r[3] or "").lower() not in _SKIP_STATUS and r[1]]
    rows = [r for r in rows if r["date"]]
    rows.sort(key=lambda r: (r["date"], r["name"].lower(), r["title"].lower()))

    wb = Workbook()

    # ── All ───────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "All"
    ws.sheet_properties.tabColor = "1E4D2F"
    first = min((r["date"] for r in rows), default=today)
    last = max((r["date"] for r in rows), default=today)
    banner(ws, "Daily Standup  |  All Tasks",
           f"{first.strftime('%d/%m/%Y')} to {last.strftime('%d/%m/%Y')}   -   one row per task per day "
           "(a task carried over shows again each day it stays open)", len(ALL_HEADERS))
    _header(ws, ALL_HEADERS, 4)
    ws.freeze_panes = "A5"
    for r in sorted(rows, key=lambda x: (x["date"], x["name"].lower()), reverse=True):
        ws.append([r["date"], r["date"].strftime("%A"), r["name"], r["title"], r["status"],
                   r["due"], r["carried_days"] or None, r["blocker"], r["subtasks"],
                   r["from"], r["to"], r["on_board"], r["added"]])
        n = ws.max_row
        ws.cell(row=n, column=1).number_format = _DATE_FMT
        if r["due"]:
            ws.cell(row=n, column=6).number_format = _DATE_FMT
        _style_status(ws, n, 5, r["status"])
        _wrap(ws, n, (4, 8))
        for c in (1, 2, 3, 5, 6, 7, 9, 10, 11, 12, 13):
            ws.cell(row=n, column=c).alignment = Alignment(
                horizontal="center" if c in (1, 2, 5, 6, 7, 9, 12) else None, vertical="center")
        _border(ws, n, len(ALL_HEADERS))
        fit_row(ws, n, [(r["title"], ALL_WIDTHS[3]), (r["blocker"], ALL_WIDTHS[7])])
    if ws.max_row > 4:
        zebra(ws, 5, ws.max_row, len(ALL_HEADERS))
        ws.auto_filter.ref = f"A4:{get_column_letter(len(ALL_HEADERS))}{ws.max_row}"
    _widths(ws, ALL_WIDTHS)
    page_setup(ws, 4)

    # ── per-person data ───────────────────────────────────────────────────
    by_uid = {}
    for r in rows:
        by_uid.setdefault(r["uid"], []).append(r)
    people = []  # (uid, name)
    seen = set()
    for e in active_employees:
        people.append((e["id"], e.get("name") or e["id"]))
        seen.add(e["id"])
    for uid in by_uid:
        if uid not in seen:
            people.append((uid, name_of(uid)))

    # ── Summary ───────────────────────────────────────────────────────────
    wsum = wb.create_sheet(title="Summary", index=1)
    wsum.sheet_properties.tabColor = "D97706"
    banner(wsum, "Daily Standup  |  Team Summary",
           f"{first.strftime('%d/%m/%Y')} to {last.strftime('%d/%m/%Y')}", 11,
           note=("Tasks are counted once each, however many days they were carried over.  "
                 "Completion % = Done / Tasks Added.  Still Open = never marked done.  "
                 "Avg Days to Finish = days from a task first appearing to being done."))
    sum_headers = ["Employee", "Active Days", "Tasks Added", "Done", "Still Open", "Blocked",
                   "Completion %", "Avg Days to Finish", "Delegated Out", "Received",
                   "Open Today"]
    _header(wsum, sum_headers, 8)
    wsum.freeze_panes = "A9"
    today_open = {}
    for r in rows:
        if r["date"] == today and r["status"] in ("Pending", "Blocked", "Needs Approval"):
            today_open[r["uid"]] = today_open.get(r["uid"], 0) + 1
    received = {}
    for r in rows:
        if r["from"]:
            received[r["name"]] = received.get(r["name"], 0) + 1
    team = {"total": 0, "done": 0, "blocked": 0}
    for uid, name in people:
        mine = by_uid.get(uid, [])
        stats = lineage_stats(mine)
        total = len(stats)
        done = sum(1 for s in stats.values() if s["done"])
        blocked = sum(1 for s in stats.values() if s["blocked"])
        finish_days = []
        for key, s in stats.items():
            if s["done"] and key[2]:
                done_dates = [r["date"] for r in mine
                              if r["lineage"] == key and r["status"] == "Done"]
                if done_dates:
                    finish_days.append((min(done_dates) - key[2]).days)
        team["total"] += total
        team["done"] += done
        team["blocked"] += blocked
        wsum.append([name, len({r["date"] for r in mine}), total, done, total - done, blocked,
                     _pct(done, total) if total else "N/A",
                     round(sum(finish_days) / len(finish_days), 1) if finish_days else "N/A",
                     sum(1 for r in mine if r["status"] == "Delegated"), received.get(name, 0),
                     today_open.get(uid, 0)])
        n = wsum.max_row
        if total:
            wsum.cell(row=n, column=7).number_format = "0.0%"
        wsum.cell(row=n, column=1).font = _BOLD
        for c in range(2, len(sum_headers) + 1):
            wsum.cell(row=n, column=c).alignment = Alignment(horizontal="center", vertical="center")
        wsum.row_dimensions[n].height = 22
        _border(wsum, n, len(sum_headers))
    if wsum.max_row >= 9:
        zebra(wsum, 9, wsum.max_row, len(sum_headers))
        wsum.auto_filter.ref = f"A8:{get_column_letter(len(sum_headers))}{wsum.max_row}"
        wsum.conditional_formatting.add(
            f"G9:G{wsum.max_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8CBCB",
                           mid_type="num", mid_value=0.5, mid_color="FFEB9C",
                           end_type="num", end_value=1, end_color="A9D8B8"))
        wsum.conditional_formatting.add(
            f"D9:D{wsum.max_row}",
            DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
    _widths(wsum, [18, 12, 12, 10, 11, 9, 14, 16, 14, 10, 11])
    # KPI cards across the top: label row 5, big number row 6
    pct_all = _pct(team["done"], team["total"])
    cards = [("TASKS ADDED", team["total"], "0", 1, 2), ("DONE", team["done"], "0", 3, 4),
             ("STILL OPEN", team["total"] - team["done"], "0", 5, 6),
             ("COMPLETION", pct_all if pct_all is not None else "N/A", "0%", 7, 8),
             ("BLOCKED", team["blocked"], "0", 9, 11)]
    for label, val, fmt, c1, c2 in cards:
        for r_ in (5, 6):
            wsum.merge_cells(start_row=r_, start_column=c1, end_row=r_, end_column=c2)
        lab = wsum.cell(row=5, column=c1, value=label)
        lab.font = Font(bold=True, size=8.5, color="6B7280")
        num = wsum.cell(row=6, column=c1, value=val)
        num.font = Font(bold=True, size=20, color="1E4D2F")
        num.number_format = fmt
        for r_ in (5, 6):
            for c_ in range(c1, c2 + 1):
                cell = wsum.cell(row=r_, column=c_)
                cell.fill = _CARD_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(left=_THIN if c_ == c1 else None,
                                     right=_THIN if c_ == c2 else None,
                                     top=_THIN if r_ == 5 else None,
                                     bottom=_THIN if r_ == 6 else None)
    wsum.row_dimensions[5].height = 18
    wsum.row_dimensions[6].height = 36
    wsum.row_dimensions[7].height = 10
    page_setup(wsum, 8)

    # ── one tracker sheet per person (one row per TASK, like the team's
    # own tracker sheets), then a combined "Collated by HR" sheet ─────────
    used = {"All", "Summary"}
    all_tracker = []
    for uid, name in people:
        mine = by_uid.get(uid, [])
        trows = tracker_rows(mine, notion_info)
        all_tracker.extend(trows)
        title = re.sub(r"[\[\]:\*\?/\\]", "", name)[:31] or str(uid)
        base, n = title, 2
        while title in used:
            title = f"{base[:28]}~{n}"
            n += 1
        used.add(title)
        w = wb.create_sheet(title=title)
        span = (f"{min(r['date'] for r in mine).strftime('%d/%m/%Y')} to "
                f"{max(r['date'] for r in mine).strftime('%d/%m/%Y')}") if mine else "no standup tasks in this range"
        banner(w, f"Daily Standup  |  {name}", span, len(TRACK_HEADERS), note=(
            "TAT = days allowed (due date - start + 1; edit freely).  Time Taken = Finish - Start + 1.  "
            "Time Track and Time Taken update by formula.  Status has a dropdown.  QR is yours to fill."))
        data_end = write_tracker(w, trows, 4)
        w.freeze_panes = "A5"
        w.sheet_properties.tabColor = "2F6B45"
        page_setup(w, 4)

        stats = lineage_stats(mine)
        months = {}
        blank = lambda: {"new": 0, "open": 0, "blocked": 0, "done": 0}  # noqa: E731
        for key, s in stats.items():
            if not s["month"]:
                continue
            m = months.setdefault(s["month"], blank())
            m["new"] += 1
            m["blocked"] += 1 if s["blocked"] else 0
            m["open"] += 0 if s["done"] else 1
        for r in mine:
            if r["status"] == "Done":
                months.setdefault((r["date"].year, r["date"].month), blank())["done"] += 1
        mh = data_end + 2
        w.cell(row=mh, column=1, value="Monthly Summary").font = _TITLE_FONT
        _header(w, ["Month", "Tasks Added", "Done", "Still Open", "Blocked", "Completion %"], mh + 1)
        k = mh + 1
        tot = blank()
        for ym in sorted(months):
            m = months[ym]
            k += 1
            pct = _pct(m["done"], m["new"])
            vals = [_date(ym[0], ym[1], 1).strftime("%B %Y"), m["new"], m["done"], m["open"],
                    m["blocked"], pct if pct is not None else "N/A"]
            for c, v in enumerate(vals, start=1):
                w.cell(row=k, column=c, value=v)
            if pct is not None:
                w.cell(row=k, column=6).number_format = "0.0%"
            _border(w, k, 6)
            for key in tot:
                tot[key] += m[key]
        k += 1
        tpct = _pct(tot["done"], tot["new"])
        for c, v in enumerate(["Total (All Time)", tot["new"], tot["done"], tot["open"],
                               tot["blocked"], tpct if tpct is not None else "N/A"], start=1):
            w.cell(row=k, column=c, value=v).font = _BOLD
            w.cell(row=k, column=c).fill = _TOTAL_FILL
        if tpct is not None:
            w.cell(row=k, column=6).number_format = "0.0%"
        _border(w, k, 6)

    wc = wb.create_sheet(title="Collated by HR", index=2)
    wc.sheet_properties.tabColor = "2F6B45"
    banner(wc, "Daily Standup  |  Collated by HR",
           "Every person's tasks, oldest first - same columns as each person's sheet; filter by Employee.",
           len(TRACK_HEADERS) + 1)
    all_tracker.sort(key=lambda t: (t["start"], t["name"].lower(), t["goal"].lower()))
    write_tracker(wc, all_tracker, 4, with_employee=True)
    wc.freeze_panes = "B5"
    page_setup(wc, 4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
